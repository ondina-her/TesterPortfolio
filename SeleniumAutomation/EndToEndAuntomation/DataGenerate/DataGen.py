import openpyxl

def dataGenerator():
    #li = [['uname1', 'lname1'],['uname2', 'lname2'],['uname3', 'lname3']]
    #return li
    wk = openpyxl.load_workbook("C:/Users/Lenovo/Documents/tests/TDataPytest.xlsx")
    sh = wk['Hoja1']
    r = sh.max_row
    li = []
    for i in range(1, r + 1):
        uname = sh.cell(row=i, column=1).value
        lname = sh.cell(row=i, column=2).value
        if uname is not None and lname is not None:
            li.append([uname, lname])

    print(li)
    return li